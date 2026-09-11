"""Test proposed cross-deposit calibration correspondence on charts, not skin TEST."""
import hashlib
import json
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/data/provenance/skin_public_2026_09_11"
he_path = next((ROOT / "data/public/he_skin_2021").glob("*.xlsx"))
liu_path = next((ROOT / "data/public/liu_spectral_2021").glob("*.xlsx"))
he = openpyxl.load_workbook(he_path, read_only=True, data_only=True)
liu = openpyxl.load_workbook(liu_path, read_only=True, data_only=True)


def cells(book, sheet, r1, r2, c1, c2):
    return np.array(list(book[sheet].iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2, values_only=True)), dtype=float)


xyz = cells(he, "CCSG", 3, 142, 3, 5)
raw_he = cells(he, "CCSG", 3, 142, 6, 8)
raw_liu = cells(liu, "RGB_SG140", 4, 143, 5, 7)
reflectance = cells(liu, "Spectral_SG140", 2, 141, 2, 32)
spd = cells(liu, "SPD_value", 3, 403, 3, 3).ravel()
cmf = np.loadtxt(OUT / "CIE_xyz_1931_2deg.csv", delimiter=",")
records = []
for name, wave in (("10nm_400_700", np.arange(400, 701, 10)), ("1nm_400_700", np.arange(400, 701)), ("1nm_380_780_constant_reflectance_tails", np.arange(380, 781))):
    illuminant = np.interp(wave, np.arange(380, 781), spd)
    observer = np.stack([np.interp(wave, cmf[:, 0], cmf[:, c]) for c in (1, 2, 3)], axis=1)
    samples = np.stack([np.interp(wave, np.arange(400, 701, 10), row) for row in reflectance])
    weights = illuminant[:, None] * observer
    weights *= 100 / weights[:, 1].sum()
    predicted = samples @ weights
    records.append({"integration": name, "white_xyz": weights.sum(axis=0).tolist(), "max_xyz_difference": float(abs(predicted-xyz).max()), "xyz_rmse": float(np.sqrt(np.mean((predicted-xyz)**2))), "tolerance_1e_minus6_pass": bool(np.allclose(predicted, xyz, rtol=0, atol=1e-6))})
result = {"scope": "Cross-deposit chart correspondence only. No FSCD/Testing numeric targets read.", "raw_exact_match": bool(np.array_equal(raw_he, raw_liu)), "raw_max_difference": float(abs(raw_he-raw_liu).max()), "records": records, "decision": "No inferred skin white point is authorized by this diagnostic alone; retain original capture/physical correspondence uncertainty.", "sha256": {"he_workbook": hashlib.sha256(he_path.read_bytes()).hexdigest(), "liu_workbook": hashlib.sha256(liu_path.read_bytes()).hexdigest(), "cmf": hashlib.sha256((OUT / "CIE_xyz_1931_2deg.csv").read_bytes()).hexdigest(), "script": hashlib.sha256(Path(__file__).read_bytes()).hexdigest()}}
destination = OUT / "he_liu_correspondence.json"
if destination.exists():
    raise ValueError("Immutable correspondence check exists")
destination.write_text(json.dumps(result, indent=2)+"\n", encoding="utf-8")
print(json.dumps(result))
